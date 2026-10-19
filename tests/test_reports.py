import importlib.util
import io
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from openpyxl import load_workbook
from reports import financial_pdf, financial_xlsx


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.horses = [{"id": "h1", "stage": "training"}]
        self.expenses = [{"id": "e1", "date": "2026-10-18", "category": "food", "amountCents": 12345, "description": "Ração"}]

    def test_xlsx_has_typed_financial_detail(self):
        payload = financial_xlsx(self.expenses, self.horses, "2026-10")
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        sheet = workbook["Relatório financeiro"]
        self.assertEqual(sheet["A1"].value, "Cocheira RADAMÉS - Relatório financeiro")
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        self.assertIn(123.45, values)

    @unittest.skipUnless(importlib.util.find_spec("reportlab"), "reportlab ainda não instalado no Python do sistema")
    def test_pdf_is_valid(self):
        payload = financial_pdf(self.expenses, self.horses, "2026-10")
        self.assertTrue(payload.startswith(b"%PDF"))
        self.assertGreater(len(payload), 1000)


if __name__ == "__main__":
    unittest.main()
