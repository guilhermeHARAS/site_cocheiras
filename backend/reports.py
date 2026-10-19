"""Relatórios financeiros baixáveis em PDF e XLSX."""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

from finance_assistant import CATEGORIES, _for_month, _period_expenses


def _period(value: str, today: date) -> tuple[dict[str, Any], str]:
    token = str(value or "").strip()
    if re.fullmatch(r"20\d{2}-(?:0[1-9]|1[0-2])", token):
        year, month = map(int, token.split("-"))
        month_name = ("Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro")[month - 1]
        return {"kind": "month", "year": year, "month": month, "label": f"{month_name} de {year}"}, token
    if re.fullmatch(r"20\d{2}", token):
        year = int(token)
        return {"kind": "year", "year": year, "label": token}, token
    raise ValueError("Período inválido para o relatório.")


def _data(expenses: list[dict[str, Any]], horses: list[dict[str, Any]], period_value: str, today: date | None = None) -> dict[str, Any]:
    current = today or date.today()
    period, token = _period(period_value, current)
    items = _period_expenses(expenses, period, current)
    items.sort(key=lambda item: (str(item.get("date") or ""), str(item.get("description") or "")))
    categories = {key: 0 for key in CATEGORIES}
    for item in items:
        key = item.get("category") if item.get("category") in CATEGORIES else "other"
        categories[key] += int(item.get("amountCents") or 0)
    total = sum(categories.values())
    fixed = sum(int(item.get("amountCents") or 0) for item in items if item.get("recurring"))
    horse_count = sum(horse.get("stage") in ("foals", "training", "ready") for horse in horses)
    return {
        "period": period,
        "token": token,
        "items": items,
        "categories": [(CATEGORIES[key], value) for key, value in categories.items() if value],
        "total": total,
        "fixed": fixed,
        "horseCount": horse_count,
        "perHorse": round(total / horse_count) if horse_count else 0,
    }


def financial_xlsx(expenses: list[dict[str, Any]], horses: list[dict[str, Any]], period_value: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    data = _data(expenses, horses, period_value)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Relatório financeiro"
    forest, gold, cream, muted = "173D2F", "C89D4A", "F6F2E9", "66736D"
    thin = Side(style="thin", color="DDD8CD")

    sheet.merge_cells("A1:F1")
    sheet["A1"] = "Cocheira RADAMÉS - Relatório financeiro"
    sheet["A1"].font = Font(name="Georgia", size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=forest)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells("A2:F2")
    sheet["A2"] = f"Período: {data['period']['label']} | Gerado em {date.today().strftime('%d/%m/%Y')}"
    sheet["A2"].font = Font(size=10, color=muted)

    metrics = [
        ("Custo total", data["total"] / 100),
        ("Despesas fixas", data["fixed"] / 100),
        ("Custo médio por cavalo", data["perHorse"] / 100),
        ("Cavalos considerados", data["horseCount"]),
    ]
    for column, (label, value) in enumerate(metrics, start=1):
        cell = sheet.cell(4, column, label)
        cell.font = Font(size=9, bold=True, color=muted)
        amount = sheet.cell(5, column, value)
        amount.font = Font(size=14, bold=True, color=forest)
        if column <= 3:
            amount.number_format = 'R$ #,##0.00'

    sheet["A7"] = "Distribuição por categoria"
    sheet["A7"].font = Font(name="Georgia", size=13, bold=True, color=forest)
    category_row = 8
    for label, value in sorted(data["categories"], key=lambda pair: pair[1], reverse=True):
        sheet.cell(category_row, 1, label)
        amount = sheet.cell(category_row, 2, value / 100)
        amount.number_format = 'R$ #,##0.00'
        category_row += 1

    detail_row = max(12, category_row + 2)
    sheet.cell(detail_row, 1, "Data")
    sheet.cell(detail_row, 2, "Categoria")
    sheet.cell(detail_row, 3, "Descrição")
    sheet.cell(detail_row, 4, "Tipo")
    sheet.cell(detail_row, 5, "Valor")
    sheet.cell(detail_row, 6, "Origem")
    for cell in sheet[detail_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=forest)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[detail_row].height = 23

    for row, item in enumerate(data["items"], start=detail_row + 1):
        try:
            sheet.cell(row, 1, datetime.strptime(str(item.get("date")), "%Y-%m-%d").date())
            sheet.cell(row, 1).number_format = "dd/mm/yyyy"
        except ValueError:
            sheet.cell(row, 1, str(item.get("date") or ""))
        sheet.cell(row, 2, CATEGORIES.get(item.get("category"), "Outros"))
        sheet.cell(row, 3, str(item.get("description") or "Despesa da cocheira"))
        sheet.cell(row, 4, "Fixa" if item.get("recurring") else "Variável")
        sheet.cell(row, 5, int(item.get("amountCents") or 0) / 100).number_format = 'R$ #,##0.00'
        sheet.cell(row, 6, "Projeção" if item.get("recurringProjection") else "Registrada")
        for cell in sheet[row]:
            cell.border = Border(bottom=thin)
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=cream)

    widths = {"A": 15, "B": 20, "C": 42, "D": 13, "E": 17, "F": 15}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = f"A{detail_row + 1}"
    sheet.auto_filter.ref = f"A{detail_row}:F{max(detail_row, detail_row + len(data['items']))}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def financial_pdf(expenses: list[dict[str, Any]], horses: list[dict[str, Any]], period_value: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    data = _data(expenses, horses, period_value)
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=17 * mm, leftMargin=17 * mm, topMargin=18 * mm, bottomMargin=17 * mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleRadames", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, textColor=colors.HexColor("#173D2F"), alignment=TA_LEFT, spaceAfter=4 * mm)
    heading = ParagraphStyle("HeadingRadames", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor("#173D2F"), spaceBefore=4 * mm, spaceAfter=2 * mm)
    body = ParagraphStyle("BodyRadames", parent=styles["BodyText"], fontSize=9, leading=12, textColor=colors.HexColor("#25362F"))
    story = [
        Paragraph("Cocheira RADAMÉS", title),
        Paragraph(f"Relatório financeiro - {data['period']['label']}", heading),
        Paragraph(f"Gerado em {date.today().strftime('%d/%m/%Y')}. Valores projetados de despesas fixas aparecem identificados na tabela.", body),
        Spacer(1, 5 * mm),
    ]
    metrics = [
        ["Custo total", "Despesas fixas", "Custo médio por cavalo", "Cavalos"],
        [_brl(data["total"]), _brl(data["fixed"]), _brl(data["perHorse"]), str(data["horseCount"])],
    ]
    metric_table = Table(metrics, colWidths=[43 * mm] * 4)
    metric_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173D2F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F6F2E9")),
        ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#173D2F")),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 7),
        ("FONTSIZE", (0, 1), (-1, 1), 10),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOX", (0, 0), (-1, -1), .5, colors.HexColor("#D9D4C9")),
        ("INNERGRID", (0, 0), (-1, -1), .25, colors.HexColor("#D9D4C9")),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([metric_table, Paragraph("Distribuição por categoria", heading)])
    category_rows = [["Categoria", "Valor"]] + [[label, _brl(value)] for label, value in sorted(data["categories"], key=lambda pair: pair[1], reverse=True)]
    story.append(_pdf_table(category_rows, [120 * mm, 52 * mm]))
    story.append(Paragraph("Lançamentos", heading))
    detail = [["Data", "Categoria", "Descrição", "Tipo", "Valor"]]
    for item in data["items"]:
        raw_date = str(item.get("date") or "")
        shown_date = f"{raw_date[8:10]}/{raw_date[5:7]}/{raw_date[:4]}" if len(raw_date) == 10 else raw_date
        kind = "Fixa projetada" if item.get("recurringProjection") else "Fixa" if item.get("recurring") else "Variável"
        detail.append([shown_date, CATEGORIES.get(item.get("category"), "Outros"), Paragraph(str(item.get("description") or "Despesa da cocheira"), body), kind, _brl(item.get("amountCents", 0))])
    story.append(_pdf_table(detail, [23 * mm, 34 * mm, 61 * mm, 28 * mm, 26 * mm]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#66736D"))
        canvas.drawString(17 * mm, 10 * mm, "Cocheira RADAMÉS - relatório somente para consulta")
        canvas.drawRightString(193 * mm, 10 * mm, f"Página {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _brl(cents: int) -> str:
    rendered = f"{int(cents or 0) / 100:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return "R$ " + rendered


def _pdf_table(rows: list[list[Any]], widths: list[float]):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    table = Table(rows, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173D2F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#D9D4C9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F2E9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (-1, 1), (-1, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table
