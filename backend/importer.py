"""Leitura segura e tolerante de planilhas e anotações para cadastro assistido."""
from __future__ import annotations

import base64
import csv
import io
import re
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from domain import STAGES

MAX_IMPORT_ROWS = 2_000
MAX_FILE_BYTES = 8 * 1024 * 1024


def _plain(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(r"[^a-z0-9]+", " ", "".join(char for char in text if not unicodedata.combining(char)).lower()).strip()


ALIASES = {
    "name": {"nome", "cavalo", "animal", "nome cavalo", "nome do cavalo", "nome animal"},
    "father": {"pai", "nome pai", "nome do pai", "garanhao", "garanhao pai", "reprodutor"},
    "mother": {"mae", "nome mae", "nome da mae", "egua", "matriz"},
    "birthDate": {"nascimento", "data nascimento", "data de nascimento", "nascido em", "dt nascimento"},
    "expectedBirthDate": {"previsao", "nascimento previsto", "data prevista", "previsao nascimento", "data prevista nascimento"},
    "age": {"idade", "idade anos", "idade em anos", "anos", "idade aproximada"},
    "breed": {"raca", "pelagem raca"},
    "sex": {"sexo", "genero"},
    "owner": {"proprietario", "dono", "socio", "titular"},
    "location": {"localizacao", "local", "baia", "piquete", "onde esta"},
    "notes": {"observacao", "observacoes", "anotacao", "anotacoes", "notas"},
    "workouts": {"treinos", "quantidade treinos", "numero treinos", "total treinos", "treinos realizados"},
    "stage": {"etapa", "fase", "situacao", "status"},
}
ALIAS_TO_FIELD = {alias: field for field, aliases in ALIASES.items() for alias in aliases}
STAGE_ALIASES = {
    "creation": "creation", "criacao": "creation", "cruza": "creation", "cruzas": "creation", "gestacao": "creation", "reproducao": "creation",
    "foals": "foals", "potro": "foals", "potros": "foals", "potra": "foals", "potras": "foals", "potros e desenvolvimento": "foals", "desenvolvimento": "foals",
    "training": "training", "treinamento": "training", "em treinamento": "training", "treino": "training", "preparacao e treinamento": "training", "preparacao": "training",
    "ready": "ready", "tropa pronta": "ready", "pronto": "ready", "pronta": "ready", "torneio": "ready",
}


def _text(value: Any, limit: int = 600) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()[:limit]


def _date_value(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 1 <= float(value) <= 100_000:
        # Datas numéricas do Excel usam 30/12/1899 como origem prática.
        return (date(1899, 12, 30) + timedelta(days=int(value))).isoformat()
    text = re.sub(r"(?i)^\s*(?:dia|em)\s+", "", _text(value, 60)).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    month_names = {
        "janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6,
        "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12,
    }
    spoken = re.fullmatch(r"\s*(\d{1,2})\s+(?:de\s+)?([a-zç]+)\s+(?:de\s+)?(\d{4})\s*", _plain(text))
    if spoken and spoken.group(2) in month_names:
        try:
            return date(int(spoken.group(3)), month_names[spoken.group(2)], int(spoken.group(1))).isoformat()
        except ValueError:
            return None
    numeric = re.fullmatch(r"\s*(\d{1,2})\s+(\d{1,2})\s+(\d{4})\s*", text)
    if numeric:
        try:
            return date(int(numeric.group(3)), int(numeric.group(2)), int(numeric.group(1))).isoformat()
        except ValueError:
            return None
    return None


def _integer(value: Any, maximum: int) -> int | None:
    if value in (None, ""):
        return None
    match = re.search(r"-?\d+", _text(value, 30))
    return max(0, min(maximum, int(match.group()))) if match else None


def _stage(value: Any, default_stage: str) -> str:
    candidate = STAGE_ALIASES.get(_plain(value), default_stage)
    return candidate if candidate in STAGES else "training"


def _record(values: dict[str, Any], default_stage: str, *, origin: str, extra: dict[str, Any] | None = None) -> dict[str, Any]:
    stage = _stage(values.get("stage"), default_stage)
    record = {
        "stage": stage,
        "name": _text(values.get("name"), 80),
        "father": _text(values.get("father"), 80),
        "mother": _text(values.get("mother"), 80),
        "birthDate": _date_value(values.get("birthDate")),
        "expectedBirthDate": _date_value(values.get("expectedBirthDate")),
        "age": _integer(values.get("age"), 60),
        "breed": _text(values.get("breed"), 80),
        "sex": _text(values.get("sex"), 20).capitalize(),
        "owner": _text(values.get("owner"), 100),
        "location": _text(values.get("location"), 120),
        "notes": _text(values.get("notes"), 600),
        "workouts": _integer(values.get("workouts"), 1_000_000) or 0,
        "importOrigin": origin,
    }
    extras = [f"{key}: {_text(value, 120)}" for key, value in (extra or {}).items() if _text(value, 120)]
    if extras:
        suffix = "Dados adicionais da importação — " + "; ".join(extras)
        record["notes"] = (record["notes"] + ("\n" if record["notes"] else "") + suffix)[:600]
    return _with_issues(record)


def _with_issues(record: dict[str, Any]) -> dict[str, Any]:
    issues: list[str] = []
    if record["stage"] == "creation":
        if not record["father"]:
            issues.append("Informe o pai")
        if not record["mother"]:
            issues.append("Informe a mãe")
        if not record["expectedBirthDate"]:
            issues.append("Informe a previsão de nascimento")
    else:
        if not record["name"]:
            issues.append("Informe o nome")
    record["issues"] = issues
    return record


def refresh_record(record: dict[str, Any], default_stage: str = "training") -> dict[str, Any]:
    """Normaliza novamente uma linha editada na tela de conferência."""
    return _record(record, default_stage, origin=_text(record.get("importOrigin"), 120) or "Cadastro rápido")


def _header_mapping(row: list[Any]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for index, cell in enumerate(row):
        field = ALIAS_TO_FIELD.get(_plain(cell))
        if field and field not in mapping.values():
            mapping[index] = field
    return mapping


def _sheet_records(rows: list[list[Any]], default_stage: str, sheet_name: str) -> tuple[list[dict[str, Any]], list[str]]:
    candidates = [(index, _header_mapping(row)) for index, row in enumerate(rows[:20])]
    header_index, mapping = max(candidates, key=lambda item: len(item[1]), default=(0, {}))
    if not mapping or "name" not in mapping.values() and not {"father", "mother"}.issubset(mapping.values()):
        return [], [f"A aba “{sheet_name}” não possui cabeçalhos reconhecíveis."]
    headers = [_text(cell, 80) or f"Coluna {index + 1}" for index, cell in enumerate(rows[header_index])]
    records: list[dict[str, Any]] = []
    warnings: list[str] = []
    pending_values: dict[str, Any] | None = None
    pending_extra: dict[str, Any] = {}
    pending_start = 0
    pending_end = 0
    merged_lines = 0

    def flush_pending() -> None:
        nonlocal pending_values, pending_extra, pending_start, pending_end
        if pending_values is None:
            return
        line_label = f"linha {pending_start}" if pending_start == pending_end else f"linhas {pending_start}–{pending_end} reunidas"
        records.append(_record(pending_values, default_stage, origin=f"{sheet_name}, {line_label}", extra=pending_extra))
        pending_values, pending_extra = None, {}

    for row_number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(_text(cell) for cell in row):
            flush_pending()
            continue
        values = {field: row[index] if index < len(row) else None for index, field in mapping.items()}
        extra = {headers[index]: row[index] for index in range(min(len(headers), len(row))) if index not in mapping}
        has_name = bool(_text(values.get("name")))
        can_continue = default_stage != "creation" and pending_values is not None and not has_name
        if can_continue:
            for field, value in values.items():
                if _text(value) and not _text(pending_values.get(field)):
                    pending_values[field] = value
            for key, value in extra.items():
                if _text(value):
                    pending_extra[key] = value
            pending_end = row_number
            merged_lines += 1
        else:
            flush_pending()
            pending_values = values
            pending_extra = extra
            pending_start = pending_end = row_number
        if len(records) >= MAX_IMPORT_ROWS:
            break
    flush_pending()
    if merged_lines:
        warnings.append(f"{merged_lines} linha(s) sem nome foram reunidas ao cavalo imediatamente anterior. Confira os registros indicados.")
    return records, warnings


def parse_spreadsheet(filename: str, content_base64: str, default_stage: str) -> dict[str, Any]:
    try:
        raw = base64.b64decode(content_base64, validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError("Não foi possível ler o arquivo enviado.") from exc
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise ValueError("A planilha deve ter no máximo 8 MB.")
    suffix = Path(filename or "").suffix.lower()
    sheets: list[tuple[str, list[list[Any]]]] = []
    if suffix == ".csv":
        text = raw.decode("utf-8-sig", errors="replace")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        sheets.append(("CSV", list(csv.reader(io.StringIO(text), dialect))[:MAX_IMPORT_ROWS + 20]))
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("O leitor de Excel ainda não está instalado no servidor.") from exc
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                rows = [list(row) for row in sheet.iter_rows(values_only=True, max_row=MAX_IMPORT_ROWS + 20)]
                sheets.append((sheet.title, rows))
        except Exception as exc:
            raise ValueError("O arquivo Excel está inválido ou protegido por senha.") from exc
    else:
        raise ValueError("Envie uma planilha .xlsx ou .csv.")

    records: list[dict[str, Any]] = []
    warnings: list[str] = []
    for name, rows in sheets:
        parsed, sheet_warnings = _sheet_records(rows, default_stage, name)
        records.extend(parsed)
        warnings.extend(sheet_warnings)
    if not records:
        raise ValueError("Nenhum cadastro foi encontrado. Use uma coluna Nome, Cavalo ou Animal.")
    if len(records) >= MAX_IMPORT_ROWS:
        warnings.append(f"Foram lidas somente as primeiras {MAX_IMPORT_ROWS} linhas.")
    return {"records": records[:MAX_IMPORT_ROWS], "warnings": warnings}


TEXT_LABELS = {
    "nome do cavalo": "name", "nome": "name", "cavalo": "name", "animal": "name",
    "nome do pai": "father", "pai": "father", "garanhao": "father",
    "nome da mae": "mother", "mae": "mother", "matriz": "mother",
    "data de nascimento": "birthDate", "nascimento": "birthDate", "nasceu": "birthDate",
    "nascimento previsto": "expectedBirthDate", "data prevista": "expectedBirthDate", "previsao": "expectedBirthDate",
    "idade aproximada": "age", "idade": "age", "raca": "breed", "sexo": "sex",
    "proprietario": "owner", "dono": "owner", "localizacao": "location", "local": "location",
    "baia": "location", "piquete": "location", "observacoes": "notes", "observacao": "notes",
    "anotacoes": "notes", "anotacao": "notes", "treinos": "workouts", "etapa": "stage", "fase": "stage",
}
# O padrão precisa reconhecer a grafia que a pessoa realmente digitou; o
# dicionário continua sendo consultado sem acentos por meio de ``_plain``.
for accented, plain in {
    "mãe": "mae", "nome da mãe": "nome da mae", "garanhão": "garanhao",
    "previsão": "previsao", "raça": "raca", "proprietário": "proprietario", "localização": "localizacao",
    "observação": "observacao", "observações": "observacoes",
    "anotação": "anotacao", "anotações": "anotacoes",
}.items():
    TEXT_LABELS[accented] = TEXT_LABELS[plain]
LABEL_PATTERN = re.compile(
    r"(?i)(?:^|[\s;,|.])(" + "|".join(sorted((re.escape(label) for label in TEXT_LABELS), key=len, reverse=True)) + r")\s*[:=\-]?\s*"
)


def _parse_text_block(block: str) -> dict[str, str]:
    matches = list(LABEL_PATTERN.finditer(block))
    values: dict[str, str] = {}
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(block)
        value = block[match.end():end].strip(" \t\r\n,;|.-")
        field = TEXT_LABELS[_plain(match.group(1))]
        if value and field not in values:
            values[field] = value
    if not values and block.strip():
        values["name"] = block.strip().splitlines()[0]
        values["notes"] = block.strip()
    return values


def _parse_simple_line(line: str) -> dict[str, Any]:
    """Interpreta listas simples como ``Relâmpago`` ou ``Relâmpago - 3 anos``."""
    clean = re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", line).strip()
    match = re.fullmatch(r"(.+?)(?:\s*[,;|\-–—]\s*|\s+)(\d{1,2})(?:\s*anos?)?\s*", clean, re.IGNORECASE)
    if match and 0 <= int(match.group(2)) <= 60:
        return {"name": match.group(1).strip(), "age": match.group(2)}
    return {"name": clean}


def _pedigree(line: str) -> tuple[str, str] | None:
    match = re.fullmatch(r"\s*(.+?)\s*[x×]\s*(.+?)\s*", line, re.IGNORECASE)
    return (match.group(1).strip(), match.group(2).strip()) if match else None


def _unstructured_block(block: str) -> list[dict[str, Any]]:
    """Lê blocos sem rótulos, usando linha em branco como separador forte."""
    lines = [re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", line).strip() for line in block.splitlines() if line.strip()]
    if not lines:
        return []

    # Linhas independentes com nome ou nome+idade continuam sendo uma lista de animais.
    secondary_information = any(
        _pedigree(line)
        or re.fullmatch(r"\d{1,2}\s*anos?", line, re.IGNORECASE)
        or _date_value(line)
        for line in lines[1:]
    )
    if len(lines) > 1 and not secondary_information:
        return [_parse_simple_line(line) for line in lines]

    first = lines[0]
    values: dict[str, Any] = {}
    warning = ""
    inline_pedigree = _pedigree(first)
    if inline_pedigree:
        left, mother = inline_pedigree
        separated = re.fullmatch(r"(.+?)\s*[;|]\s*(.+)", left)
        if separated:
            values.update({"name": separated.group(1).strip(), "father": separated.group(2).strip(), "mother": mother})
        else:
            pieces = left.split()
            if len(pieces) >= 2:
                values.update({"name": pieces[0], "father": " ".join(pieces[1:]), "mother": mother})
                warning = "Confira a interpretação automática de nome, pai e mãe"
            else:
                values["name"] = first
    else:
        values.update(_parse_simple_line(first))

    notes: list[str] = []
    for line in lines[1:]:
        parents = _pedigree(line)
        if parents and not values.get("father") and not values.get("mother"):
            values["father"], values["mother"] = parents
            continue
        age = re.fullmatch(r"(\d{1,2})\s*anos?", line, re.IGNORECASE)
        if age and values.get("age") in (None, ""):
            values["age"] = age.group(1)
            continue
        parsed_date = _date_value(line)
        if parsed_date and not values.get("birthDate"):
            values["birthDate"] = parsed_date
            continue
        notes.append(line)
    if notes:
        values["notes"] = "\n".join(notes)
    if warning:
        values["_warning"] = warning
    return [values]


def parse_notes(text: str, default_stage: str, origin: str = "Anotações") -> dict[str, Any]:
    clean = str(text or "").replace("\r\n", "\n").strip()
    if not clean:
        raise ValueError("Cole ou dite alguma informação antes de continuar.")
    if LABEL_PATTERN.search(clean):
        # Um novo Nome/Cavalo no início da linha inicia outro cadastro, mesmo sem linha em branco.
        clean = re.sub(r"(?im)(?<!\A)^(?=\s*(?:nome(?: do cavalo)?|cavalo|animal)\s*[:=\-])", "\n\n", clean)
        blocks = [part.strip() for part in re.split(r"\n\s*\n|\n\s*-{3,}\s*\n", clean) if part.strip()]
        values = [_parse_text_block(block) for block in blocks]
    else:
        blocks = [part.strip() for part in re.split(r"\n\s*\n|\n\s*-{3,}\s*\n", clean) if part.strip()]
        values = []
        for block in blocks:
            values.extend(_unstructured_block(block))
    records = []
    for index, value in enumerate(values, 1):
        if not value.get("name"):
            continue
        record = _record(value, default_stage, origin=f"{origin}, bloco {index}")
        if value.get("_warning"):
            record["issues"].append(value["_warning"])
        records.append(record)
    if not records:
        raise ValueError("Nenhum nome de cavalo foi encontrado.")
    return {"records": records[:MAX_IMPORT_ROWS], "warnings": []}


def duplicate_of(record: dict[str, Any], horses: list[dict[str, Any]]) -> str | None:
    name = _plain(record.get("name"))
    if not name:
        return None
    for horse in horses:
        if _plain(horse.get("name")) == name:
            return _text(horse.get("name"), 80)
    return None
